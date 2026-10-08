"""Read-only Excel importer with loss-aware SQLite persistence."""

from __future__ import annotations

import hashlib
import json
import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app import __version__
from app.models.roadmap import ImportIssue, ImportResult, RoadmapItem
from app.importers.validation import issue_to_dict, validate_workbook


def _json_value(value: Any) -> Any:
    """Convert openpyxl values to JSON-safe values without losing meaning."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return {"type": type(value).__name__, "value": value.isoformat()}
    return {"type": type(value).__name__, "value": str(value)}


def _dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


def _chart_title(chart: Any) -> str:
    """Extract a chart title when the workbook provides literal title text."""
    try:
        paragraphs = chart.title.tx.rich.p
        return " ".join(run.text for paragraph in paragraphs for run in paragraph.r if run.text)
    except (AttributeError, TypeError):
        return ""


def _chart_series(series: Any) -> dict[str, str | None]:
    """Persist source cell references used by a chart series."""
    title = None
    try:
        title = series.tx.strRef.f
    except AttributeError:
        pass
    values = None
    try:
        values = series.val.numRef.f
    except AttributeError:
        pass
    categories = None
    try:
        categories = series.cat.strRef.f
    except AttributeError:
        try:
            categories = series.cat.numRef.f
        except AttributeError:
            pass
    return {"title_ref": title, "values_ref": values, "categories_ref": categories}


def _row_values(ws: Worksheet, row_number: int) -> dict[str, Any]:
    return {
        get_column_letter(col): _json_value(ws.cell(row_number, col).value)
        for col in range(1, ws.max_column + 1)
        if ws.cell(row_number, col).value is not None
    }


def _item(
    ws: Worksheet,
    row_number: int,
    item_type: str,
    name: Any,
    description: Any = "",
    category: Any = "",
    phase: Any = "",
    target_period: Any = "",
    status: Any = "",
    priority: Any = "",
    xp_reward: Any = None,
    xp_earned: Any = None,
    resources: tuple[Any, ...] = (),
    notes: Any = "",
    prerequisite_notes: Any = "",
    extra: dict[str, Any] | None = None,
) -> RoadmapItem | None:
    if name in (None, ""):
        return None
    clean = lambda value: "" if value is None else str(value).strip()
    numeric_reward = int(xp_reward) if isinstance(xp_reward, (int, float)) else None
    numeric_earned = int(xp_earned) if isinstance(xp_earned, (int, float)) else None
    resource_values = tuple(str(value).strip() for value in resources if value not in (None, ""))
    source = _row_values(ws, row_number)
    if extra:
        source["_section"] = extra.get("section", "")
        source["_identity"] = extra.get("identity", "")
    return RoadmapItem(
        item_id=f"{ws.title}:{row_number}",
        source_sheet=ws.title,
        source_row=row_number,
        item_type=item_type,
        name=clean(name),
        description=clean(description),
        category=clean(category),
        phase=clean(phase),
        target_period=clean(target_period),
        status=clean(status),
        priority=clean(priority),
        xp_reward=numeric_reward,
        xp_earned=numeric_earned,
        prerequisite_notes=clean(prerequisite_notes),
        resources=resource_values,
        notes=clean(notes),
        source_record=source,
    )


def normalize_records(workbook: Any, cached_workbook: Any | None = None) -> list[RoadmapItem]:
    """Map task-like rows into a shared model while retaining all raw rows."""
    items: list[RoadmapItem] = []

    ws = workbook["Roadmap"]
    for row in range(2, 64):
        reward = 100  # The sheet's saved formula and Dashboard instructions use 100 XP for Done.
        earned = (
            cached_workbook["Roadmap"].cell(row, 8).value
            if cached_workbook is not None else ws.cell(row, 8).value
        )
        if not isinstance(earned, (int, float)):
            earned = None
        candidate = _item(
            ws, row, "monthly_mission", ws.cell(row, 5).value,
            description=ws.cell(row, 6).value, category=ws.cell(row, 4).value,
            phase=ws.cell(row, 3).value, target_period=ws.cell(row, 1).value,
            status=ws.cell(row, 7).value, xp_reward=reward, xp_earned=earned,
            notes=ws.cell(row, 10).value,
        )
        if candidate:
            items.append(candidate)

    ws = workbook["Skill Matrix"]
    for row in range(2, 33):
        candidate = _item(
            ws, row, "skill", ws.cell(row, 1).value,
            description=ws.cell(row, 4).value, category=ws.cell(row, 2).value,
            target_period=ws.cell(row, 3).value, status=ws.cell(row, 5).value,
            xp_reward=ws.cell(row, 6).value,
        )
        if candidate:
            items.append(candidate)

    ws = workbook["Certifications"]
    for row in range(2, 21):
        target_window = ws.cell(row, 5).value
        prerequisite = (
            target_window if isinstance(target_window, str)
            and any(term in target_window.casefold() for term in ("after ", "before ")) else ""
        )
        candidate = _item(
            ws, row, "certification", ws.cell(row, 1).value,
            category=ws.cell(row, 3).value, target_period=target_window or ws.cell(row, 4).value,
            status=ws.cell(row, 2).value, priority=ws.cell(row, 6).value,
            prerequisite_notes=prerequisite,
            notes=f"Target age: {ws.cell(row, 4).value or ''}".strip(),
        )
        if candidate:
            items.append(candidate)

    ws = workbook["Projects"]
    for row in range(2, 19):
        candidate = _item(
            ws, row, "project", ws.cell(row, 2).value,
            description=ws.cell(row, 4).value, category=ws.cell(row, 3).value,
            phase=ws.cell(row, 5).value, target_period=ws.cell(row, 5).value,
            status=ws.cell(row, 7).value, xp_reward=ws.cell(row, 6).value,
            extra={"identity": ws.cell(row, 1).value},
        )
        if candidate:
            items.append(candidate)

    # 12-Month Starter contains three distinct sections, sharing the same headers.
    ws = workbook["12-Month Starter"]
    for first, last, section in ((8, 53, "12-month core"), (60, 71, "Microsoft / RDP add-on"), (74, 76, "CCNA mini-track")):
        for row in range(first, last + 1):
            candidate = _item(
                ws, row, "learning_objective", ws.cell(row, 4).value,
                description=f"What to learn: {ws.cell(row, 5).value or ''}\nPractical mission: {ws.cell(row, 6).value or ''}".strip(),
                category=ws.cell(row, 3).value, phase=section,
                target_period=ws.cell(row, 1).value,
                status=ws.cell(row, 11).value,
                priority=ws.cell(row, 9).value,
                xp_reward=ws.cell(row, 10).value,
                resources=(ws.cell(row, 7).value, ws.cell(row, 8).value),
                prerequisite_notes=ws.cell(row, 2).value if section == "CCNA mini-track" else "",
                extra={"section": section},
            )
            if candidate:
                items.append(candidate)

    ws = workbook["Microsoft Track"]
    for row in range(9, 41):
        candidate = _item(
            ws, row, "learning_objective", ws.cell(row, 4).value,
            description=f"What to learn: {ws.cell(row, 5).value or ''}\nPractical mission: {ws.cell(row, 6).value or ''}".strip(),
            category=ws.cell(row, 3).value, phase=ws.cell(row, 2).value,
            target_period=ws.cell(row, 1).value, status=ws.cell(row, 11).value,
            xp_reward=ws.cell(row, 10).value,
            resources=(ws.cell(row, 7).value, ws.cell(row, 8).value),
            notes=ws.cell(row, 12).value,
            extra={"section": "Microsoft mission", "identity": ws.cell(row, 9).value},
        )
        if candidate:
            items.append(candidate)
    for row in range(44, 52):
        candidate = _item(
            ws, row, "certification", ws.cell(row, 2).value,
            description=f"{ws.cell(row, 3).value or ''}\nFocus: {ws.cell(row, 7).value or ''}".strip(),
            category="Microsoft certification ladder", target_period=ws.cell(row, 5).value,
            status=ws.cell(row, 6).value, priority=ws.cell(row, 4).value,
            extra={"section": "recommended order", "identity": ws.cell(row, 1).value},
        )
        if candidate:
            items.append(candidate)

    ws = workbook["My Certifications"]
    for row in range(5, 29):
        name = ws.cell(row, 1).value
        if name in (None, "") or str(name).strip().casefold() == "custom future certification":
            continue
        candidate = _item(
            ws, row, "certification", name,
            category=ws.cell(row, 3).value, target_period=ws.cell(row, 7).value or ws.cell(row, 4).value,
            status=ws.cell(row, 6).value, priority=ws.cell(row, 5).value,
            xp_reward=ws.cell(row, 9).value,
            resources=(ws.cell(row, 10).value,), notes=ws.cell(row, 11).value,
            prerequisite_notes=(
                ws.cell(row, 11).value
                if isinstance(ws.cell(row, 11).value, str)
                and "after " in ws.cell(row, 11).value.casefold() else ""
            ),
            extra={"section": "personal certification wallet", "identity": ws.cell(row, 2).value},
        )
        if candidate:
            items.append(candidate)

    ws = workbook["Scripting & Dev"]
    for row in range(5, 38):
        tool, skill = ws.cell(row, 3).value, ws.cell(row, 4).value
        name = f"{tool} — {skill}" if tool and skill else skill or tool
        candidate = _item(
            ws, row, "learning_objective", name,
            description=ws.cell(row, 5).value, category=ws.cell(row, 7).value,
            phase=ws.cell(row, 1).value, target_period=ws.cell(row, 2).value,
            status=ws.cell(row, 8).value, xp_reward=ws.cell(row, 9).value,
            resources=(ws.cell(row, 6).value,), notes=ws.cell(row, 10).value,
            prerequisite_notes=(
                ws.cell(row, 10).value
                if isinstance(ws.cell(row, 10).value, str)
                and "after " in ws.cell(row, 10).value.casefold() else ""
            ),
        )
        if candidate:
            items.append(candidate)

    ws = workbook["Beyond 25"]
    for row in range(5, 13):
        candidate = _item(
            ws, row, "specialization", ws.cell(row, 4).value,
            description=f"Focus: {ws.cell(row, 3).value or ''}\nPractical mission: {ws.cell(row, 5).value or ''}".strip(),
            category=ws.cell(row, 2).value, target_period=ws.cell(row, 1).value,
            status=ws.cell(row, 9).value, priority=ws.cell(row, 8).value,
            xp_reward=ws.cell(row, 10).value,
            resources=(ws.cell(row, 6).value, ws.cell(row, 7).value),
        )
        if candidate:
            items.append(candidate)

    return items


def import_excel(connection: Any, source_path: str | Path) -> ImportResult:
    """Import a workbook without modifying it; repeated identical imports are idempotent."""
    path = Path(source_path).expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Select an Excel .xlsx or .xlsm workbook.")

    hasher = hashlib.sha256()
    with path.open("rb") as source_file:
        for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
            hasher.update(chunk)
    workbook_id = hasher.hexdigest()
    existing = connection.execute(
        "SELECT * FROM workbooks WHERE workbook_id = ?", (workbook_id,)
    ).fetchone()
    if existing:
        with connection:
            connection.execute(
                "INSERT OR REPLACE INTO user_settings(setting_key, value_json) VALUES('active_workbook_id', ?)",
                (_dump(workbook_id),),
            )
        old_issues = connection.execute(
            "SELECT severity, code, message, sheet_name, cell_reference FROM import_issues WHERE workbook_id = ? ORDER BY issue_id",
            (workbook_id,),
        ).fetchall()
        return ImportResult(
            workbook_id=workbook_id,
            filename=existing["filename"],
            sheet_count=existing["sheet_count"],
            source_row_count=existing["source_row_count"],
            source_cell_count=existing["source_cell_count"],
            formula_count=existing["formula_count"],
            item_count=existing["item_count"],
            issues=tuple(ImportIssue(r["severity"], r["code"], r["message"], r["sheet_name"], r["cell_reference"]) for r in old_issues),
            already_imported=True,
        )

    with warnings.catch_warnings(record=True) as captured_warnings:
        warnings.simplefilter("always")
        workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
        cached_workbook = load_workbook(path, data_only=True, read_only=False, keep_links=True)

    items = normalize_records(workbook, cached_workbook)
    cached_sheets = {ws.title: cached_workbook[ws.title] for ws in cached_workbook.worksheets}
    issues = validate_workbook(workbook.worksheets, cached_sheets, items)
    extension_warnings = [
        str(warning.message) for warning in captured_warnings
        if "extension is not supported" in str(warning.message).casefold()
    ]
    if extension_warnings:
        issues.append(ImportIssue(
            severity="info",
            code="EXCEL_EXTENSION_METADATA",
            message="The workbook contains Excel extension records for formatting or validation that openpyxl does not fully model. Cell values, formulas, cached values, standard validation ranges, and chart references are imported; the source workbook was not saved or modified.",
        ))
    formula_count = 0
    source_row_count = 0
    source_cell_count = 0
    sheet_counts: dict[str, dict[str, int]] = {}

    with connection:
        connection.execute(
            """INSERT INTO workbooks(workbook_id, filename, importer_version, sheet_count,
               source_row_count, source_cell_count, formula_count, item_count, summary_json)
               VALUES(?, ?, ?, ?, 0, 0, 0, ?, '{}')""",
            (workbook_id, path.name, __version__, len(workbook.worksheets), len(items)),
        )

        for position, ws in enumerate(workbook.worksheets, start=1):
            cached_ws = cached_workbook[ws.title]
            validations = []
            for rule in ws.data_validations.dataValidation:
                validations.append({
                    "ranges": str(rule.sqref),
                    "type": rule.type,
                    "formula1": rule.formula1,
                    "formula2": rule.formula2,
                })
            metadata = {
                "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
                "tables": [
                    {"name": name, "ref": table.ref if hasattr(table, "ref") else str(table)}
                    for name, table in ws.tables.items()
                ],
                "data_validations": validations,
                "conditional_formatting_rule_count": len(ws.conditional_formatting),
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "auto_filter": ws.auto_filter.ref,
                "show_grid_lines": ws.sheet_view.showGridLines,
                "charts": [
                    {
                        "type": type(chart).__name__,
                        "title": _chart_title(chart),
                        "series": [_chart_series(series) for series in chart.series],
                    }
                    for chart in ws._charts
                ],
            }
            connection.execute(
                """INSERT INTO source_sheets(workbook_id, sheet_name, position, state,
                   max_row, max_column, metadata_json) VALUES(?, ?, ?, ?, ?, ?, ?)""",
                (workbook_id, ws.title, position, ws.sheet_state, ws.max_row, ws.max_column, _dump(metadata)),
            )
            row_count = 0
            cell_count = 0
            for row in ws.iter_rows():
                row_number = row[0].row if row else 0
                payload: dict[str, Any] = {}
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    formula_text = value if isinstance(value, str) and value.startswith("=") else None
                    if formula_text:
                        formula_count += 1
                    cached_value = cached_ws[cell.coordinate].value if formula_text else None
                    payload[cell.column_letter] = {
                        "value": _json_value(value),
                        "formula": formula_text,
                        "cached_value": _json_value(cached_value),
                    }
                    connection.execute(
                        """INSERT INTO source_cells(workbook_id, sheet_name, coordinate,
                           row_number, column_number, data_type, value_json, formula_text,
                           cached_value_json, number_format, hyperlink)
                           VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (
                            workbook_id, ws.title, cell.coordinate, cell.row, cell.column,
                            cell.data_type, _dump(_json_value(value)), formula_text,
                            _dump(_json_value(cached_value)) if formula_text else None,
                            cell.number_format, cell.hyperlink.target if cell.hyperlink else None,
                        ),
                    )
                    cell_count += 1
                    source_cell_count += 1
                if payload:
                    connection.execute(
                        "INSERT INTO source_rows(workbook_id, sheet_name, row_number, values_json) VALUES(?, ?, ?, ?)",
                        (workbook_id, ws.title, row_number, _dump(payload)),
                    )
                    row_count += 1
                    source_row_count += 1
            sheet_counts[ws.title] = {"rows": row_count, "cells": cell_count}

        for item in items:
            connection.execute(
                """INSERT INTO roadmap_items(
                   item_id, workbook_id, source_sheet, source_row, item_type, name,
                   description, category, phase, target_period, status, priority,
                   xp_reward, xp_earned, prerequisite_notes, resources_json, notes,
                   source_record_json)
                   VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    f"{workbook_id}:{item.item_id}", workbook_id, item.source_sheet, item.source_row,
                    item.item_type, item.name, item.description, item.category,
                    item.phase, item.target_period, item.status, item.priority,
                    item.xp_reward, item.xp_earned, item.prerequisite_notes,
                    _dump(item.resources), item.notes, _dump(item.source_record),
                ),
            )

        for issue in issues:
            connection.execute(
                """INSERT INTO import_issues(workbook_id, severity, code,
                   sheet_name, cell_reference, message) VALUES(?, ?, ?, ?, ?, ?)""",
                (workbook_id, issue.severity, issue.code, issue.sheet, issue.cell, issue.message),
            )

        summary = {
            "sheet_counts": sheet_counts,
            "profile_name": workbook["Dashboard"]["B4"].value if "Dashboard" in workbook else None,
            "core_roadmap_months": sum(1 for item in items if item.item_type == "monthly_mission"),
            "issue_count": len(issues),
        }
        connection.execute(
            "UPDATE workbooks SET source_row_count=?, source_cell_count=?, formula_count=?, summary_json=? WHERE workbook_id=?",
            (source_row_count, source_cell_count, formula_count, _dump(summary), workbook_id),
        )
        connection.execute(
            "INSERT OR REPLACE INTO user_settings(setting_key, value_json) VALUES('active_workbook_id', ?)",
            (_dump(workbook_id),),
        )

    return ImportResult(
        workbook_id=workbook_id,
        filename=path.name,
        sheet_count=len(workbook.worksheets),
        source_row_count=source_row_count,
        source_cell_count=source_cell_count,
        formula_count=formula_count,
        item_count=len(items),
        issues=tuple(issues),
    )
