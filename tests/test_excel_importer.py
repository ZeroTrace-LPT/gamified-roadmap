"""Tests for source-preserving Excel import."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.database.connection import connect_database
from app.importers.excel_importer import import_excel


class ExcelImporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1])
        self.root = Path(self.temp_dir.name)
        self.workbook_path = self.root / "source.xlsx"
        self.database_path = self.root / "roadmap.sqlite3"
        self._make_workbook()
        self.connection = connect_database(self.database_path)

    def tearDown(self) -> None:
        self.connection.close()
        self.temp_dir.cleanup()

    def _make_workbook(self) -> None:
        workbook = Workbook()
        dashboard = workbook.active
        dashboard.title = "Dashboard"
        dashboard["B4"] = "Local Test Profile"
        dashboard["E8"] = "=SUM(Roadmap!H2:H63)"
        dashboard["E9"] = 1
        dashboard["C15"] = "IT Apprentice"

        roadmap = workbook.create_sheet("Roadmap")
        roadmap.append(["Month", "Age", "Year", "Track", "Focus", "Monthly Mission", "Status", "XP", "Done?", "Notes"])
        roadmap.append(["2026-09", "19", "Year 1", "Foundation", "Network+ / Linux / Git", "Build a lab", "Not Started", 0, "No", "source note"])

        skills = workbook.create_sheet("Skill Matrix")
        skills.append(["Skill", "Track", "Target Age", "What Good Looks Like", "Status", "XP Value"])
        skills.append(["Networking", "Foundation", "20", "TCP/IP", "Not Started", 500])
        certs = workbook.create_sheet("Certifications")
        certs.append(["Certification", "Status", "Domain", "Target Age", "Target Window", "Priority"])
        certs.append(["CompTIA A+", "Completed", "IT Foundations", "19", "Completed", "Done"])
        projects = workbook.create_sheet("Projects")
        projects.append(["ID", "Project", "Track", "Definition of Done", "Target", "XP", "Status"])
        projects.append(["P01", "Homelab v1", "Foundation", "AD + Windows + Linux", "Year 1", 100, "Not Started"])

        starter = workbook.create_sheet("12-Month Starter")
        starter["C5"] = "=IF(B5=0,0,C5/B5)"
        workbook.create_sheet("Microsoft Track")
        wallet = workbook.create_sheet("My Certifications")
        wallet["A25"] = "Cisco CCNA"
        wallet["F25"] = "Planned"
        workbook.create_sheet("Scripting & Dev")
        workbook.create_sheet("Beyond 25")
        workbook.create_sheet("Weekly Tracker")
        workbook.create_sheet("GitHub Portfolio")
        workbook.create_sheet("Resources")
        workbook.create_sheet("Lists")
        workbook.save(self.workbook_path)

    def test_import_persists_source_rows_cells_and_formulas(self) -> None:
        result = import_excel(self.connection, self.workbook_path)

        self.assertEqual(result.sheet_count, 14)
        self.assertGreater(result.source_cell_count, 20)
        self.assertEqual(result.formula_count, 2)
        formula = self.connection.execute(
            "SELECT formula_text FROM source_cells WHERE sheet_name='12-Month Starter' AND coordinate='C5'"
        ).fetchone()[0]
        self.assertEqual(formula, "=IF(B5=0,0,C5/B5)")
        source_note = self.connection.execute(
            "SELECT values_json FROM source_rows WHERE sheet_name='Roadmap' AND row_number=2"
        ).fetchone()[0]
        self.assertIn("source note", source_note)
        self.assertTrue(any(issue.code == "FORMULA_CIRCULAR_REFERENCE" for issue in result.issues))

    def test_normalized_item_keeps_status_and_source_reward(self) -> None:
        import_excel(self.connection, self.workbook_path)
        row = self.connection.execute(
            "SELECT name, status, xp_reward, category, target_period FROM roadmap_items WHERE item_type='monthly_mission'"
        ).fetchone()
        self.assertEqual(tuple(row), ("Network+ / Linux / Git", "Not Started", 100, "Foundation", "2026-09"))

    def test_reimport_of_identical_workbook_is_idempotent(self) -> None:
        first = import_excel(self.connection, self.workbook_path)
        second = import_excel(self.connection, self.workbook_path)
        self.assertFalse(first.already_imported)
        self.assertTrue(second.already_imported)
        self.assertEqual(self.connection.execute("SELECT COUNT(*) FROM workbooks").fetchone()[0], 1)
        self.assertEqual(self.connection.execute("SELECT COUNT(*) FROM roadmap_items").fetchone()[0], 5)

    def test_changed_workbook_creates_a_distinct_snapshot(self) -> None:
        first = import_excel(self.connection, self.workbook_path)
        workbook = load_workbook(self.workbook_path)
        workbook["Roadmap"]["A2"] = "2026-10"
        workbook.save(self.workbook_path)

        second = import_excel(self.connection, self.workbook_path)
        self.assertNotEqual(first.workbook_id, second.workbook_id)
        self.assertEqual(self.connection.execute("SELECT COUNT(*) FROM workbooks").fetchone()[0], 2)
        self.assertEqual(self.connection.execute("SELECT COUNT(*) FROM roadmap_items").fetchone()[0], 10)


if __name__ == "__main__":
    unittest.main()
